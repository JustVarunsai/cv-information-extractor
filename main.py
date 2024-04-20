import os
import re
import PyPDF2
import docx
from docx import Document
from openpyxl import Workbook


def extract_text_from_pdf(pdf_file):
    text = ""
    with open(pdf_file, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        for page_num in range(len(pdf_reader.pages)):
            page = pdf_reader.pages[page_num]
            text += page.extract_text()
    return text


def extract_text_from_docx(docx_file):
    text = ""
    doc = Document(docx_file)
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text


def extract_email(text):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    emails = re.findall(email_pattern, text)
    return emails


def extract_phone_number(text):
    phone_pattern = r'\b(?:0|\+?91)?[789]\d{9}\b'
    phone_numbers = re.findall(phone_pattern, text)
    return phone_numbers


def process_cv(cv_dir):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Email", "Phone", "Text"])

    for filename in os.listdir(cv_dir):
        if filename.endswith(".pdf"):
            text = extract_text_from_pdf(os.path.join(cv_dir, filename))
        elif filename.endswith(".docx"):
            text = extract_text_from_docx(os.path.join(cv_dir, filename))
        else:
            continue

        emails = extract_email(text)
        phone_numbers = extract_phone_number(text)

        ws.append([filename, ", ".join(emails), ", ".join(phone_numbers), text])

    wb.save("cv_info.xlsx")


if __name__ == "__main__":
    process_cv("Sample2")